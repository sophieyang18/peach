from __future__ import annotations

import io
import re
import zipfile
from dataclasses import dataclass
from html import escape, unescape
from pathlib import Path
from urllib.parse import urlparse

import httpx
from bs4 import BeautifulSoup
from docx import Document
from pypdf import PdfReader

try:
    import olefile
except ImportError:  # pragma: no cover - optional parser enhancement
    olefile = None


SUPPORTED_EXTENSIONS = {".pdf", ".doc", ".docx", ".md", ".markdown", ".html", ".htm"}
MAX_LINK_BYTES = 2 * 1024 * 1024


@dataclass
class ParsedFile:
    filename: str
    extension: str
    title: str
    content: str
    summary: str
    content_html: str = ""
    warning: str = ""


def parse_upload(filename: str, content: bytes) -> ParsedFile:
    extension = Path(filename).suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError("unsupported file type")

    warning = ""
    content_html = ""
    if extension == ".pdf":
        text = parse_pdf(content)
        content_html = parse_pdf_html(content)
    elif extension == ".docx":
        text = parse_docx(content)
        content_html = parse_docx_html(content)
    elif extension == ".doc":
        text = parse_doc_binary_best_effort(content)
        warning = "老版 doc 文件采用基础文本抽取，复杂排版可能不完整。"
    elif extension in {".html", ".htm"}:
        text = parse_html(content)
        content_html = sanitize_html_for_editor(content)
    else:
        text = parse_markdown(content)

    title = file_title(filename, text)
    clean = normalize_text(text)
    return ParsedFile(
        filename=filename,
        extension=extension.lstrip("."),
        title=title,
        content=clean,
        summary=summarize(clean),
        content_html=content_html or plain_text_to_html(clean),
        warning=warning,
    )


def parse_pdf(content: bytes) -> str:
    reader = PdfReader(io.BytesIO(content))
    pages = []
    for page in reader.pages:
        pages.append(page.extract_text() or "")
    return "\n\n".join(pages)


def parse_pdf_html(content: bytes) -> str:
    reader = PdfReader(io.BytesIO(content))
    pages: list[str] = []
    for page in reader.pages[:20]:
        fragments: list[str] = []

        def visitor_text(text, _cm, _tm, font_dict, font_size):
            if not text:
                return
            styles = []
            if font_size:
                styles.append(f"font-size:{max(8, min(24, round(float(font_size), 1)))}pt")
            base_font = ""
            if isinstance(font_dict, dict):
                base_font = str(font_dict.get("/BaseFont") or "").split("+")[-1].replace("-", " ")
            if base_font:
                styles.append(f"font-family:{escape(base_font, quote=True)}")
            style_attr = f" style=\"{';'.join(styles)}\"" if styles else ""
            fragments.append(f"<span{style_attr}>{escape(text).replace(chr(10), '<br />')}</span>")

        try:
            page.extract_text(visitor_text=visitor_text)
        except TypeError:
            text = page.extract_text() or ""
            fragments.append(plain_text_to_html(text))
        page_html = "".join(fragments).strip()
        if page_html:
            pages.append(f"<section>{page_html}</section>")
    return "\n".join(pages)


def parse_docx(content: bytes) -> str:
    document = Document(io.BytesIO(content))
    parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def parse_docx_html(content: bytes) -> str:
    document = Document(io.BytesIO(content))
    parts: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            parts.append("<p><br /></p>")
            continue
        tag = "h3" if (paragraph.style and "heading" in paragraph.style.name.lower()) or len(text) <= 16 and re.search(r"(教育背景|实习经历|项目经历|个人技能|竞赛经历|获奖经历|荣誉经历|求职意向)", text) else "p"
        runs = []
        for run in paragraph.runs:
            if not run.text:
                continue
            styles = []
            if run.bold:
                styles.append("font-weight:700")
            if run.italic:
                styles.append("font-style:italic")
            if run.font.size:
                styles.append(f"font-size:{max(9, min(24, round(run.font.size.pt)))}pt")
            if run.font.name:
                styles.append(f"font-family:{escape(run.font.name, quote=True)}")
            style_attr = f" style=\"{';'.join(styles)}\"" if styles else ""
            runs.append(f"<span{style_attr}>{escape(run.text)}</span>")
        parts.append(f"<{tag}>{''.join(runs) or escape(text)}</{tag}>")
    for table in document.tables:
        rows = []
        for row in table.rows:
            cells = "".join(f"<td>{plain_text_to_html(cell.text)}</td>" for cell in row.cells)
            if cells:
                rows.append(f"<tr>{cells}</tr>")
        if rows:
            parts.append(f"<table>{''.join(rows)}</table>")
    return "\n".join(parts)


def parse_doc_binary_best_effort(content: bytes) -> str:
    if zipfile.is_zipfile(io.BytesIO(content)):
        return parse_docx(content)

    streams = doc_stream_candidates(content)
    candidates = [extract_readable_text(stream) for stream in streams]
    candidates.append(extract_readable_text(content))
    best = max(candidates, key=text_quality_score, default="")
    best = normalize_text(best)
    if text_quality_score(best) < 18 or is_probably_garbled(best):
        raise ValueError("老版 .doc 文本抽取失败。建议将文件另存为 .docx 后重新上传。")
    return best


def parse_html(content: bytes) -> str:
    raw = decode_text_best_effort(content)
    soup = BeautifulSoup(raw, "html.parser")
    for tag in soup(["script", "style", "noscript"]):
        tag.decompose()
    return soup.get_text("\n")


def sanitize_html_for_editor(content: bytes) -> str:
    raw = decode_text_best_effort(content)
    soup = BeautifulSoup(raw, "html.parser")
    for tag in soup(["script", "style", "iframe", "object", "embed"]):
        tag.decompose()
    allowed = {"h1", "h2", "h3", "p", "ul", "ol", "li", "table", "tr", "td", "th", "strong", "b", "em", "i", "span", "br"}
    for tag in soup.find_all(True):
        if tag.name not in allowed:
            tag.unwrap()
            continue
        safe_style = tag.get("style", "")
        tag.attrs = {}
        if tag.name == "span" and safe_style:
            whitelisted = []
            for rule in safe_style.split(";"):
                name, _, value = rule.partition(":")
                if name.strip().lower() in {"font-size", "font-family", "font-weight", "font-style"}:
                    whitelisted.append(f"{name.strip()}:{escape(value.strip(), quote=True)}")
            if whitelisted:
                tag["style"] = ";".join(whitelisted)
    body = soup.body or soup
    return "\n".join(str(tag) for tag in body.find_all(list(allowed), recursive=True)[:240])


async def parse_link(url: str) -> ParsedFile:
    parsed_url = urlparse(url)
    if parsed_url.scheme not in {"http", "https"} or not parsed_url.netloc:
        raise ValueError("unsupported url")

    async with httpx.AsyncClient(timeout=12, follow_redirects=True) as client:
        response = await client.get(url, headers={"User-Agent": "PeachAgent/0.1"})
        response.raise_for_status()
        content_type = response.headers.get("content-type", "")
        content = response.content[:MAX_LINK_BYTES]

    if "pdf" in content_type or parsed_url.path.lower().endswith(".pdf"):
        parsed = parse_upload(Path(parsed_url.path).name or "link.pdf", content)
        return ParsedFile(
            filename=Path(parsed_url.path).name or parsed.filename,
            extension=parsed.extension,
            title=parsed.title,
            content=parsed.content,
            summary=parsed.summary,
            warning=parsed.warning,
        )

    text = parse_html(content)
    clean = normalize_text(text)
    title = title_from_html(content) or parsed_url.netloc.replace("www.", "")
    return ParsedFile(
        filename=parsed_url.netloc,
        extension="html",
        title=title,
        content=f"来源链接：{url}\n\n{clean}",
        summary=summarize(clean),
    )


def parse_markdown(content: bytes) -> str:
    raw = decode_text_best_effort(content)
    raw = re.sub(r"```[\s\S]*?```", " ", raw)
    raw = re.sub(r"`([^`]+)`", r"\1", raw)
    raw = re.sub(r"!\[[^\]]*\]\([^)]+\)", " ", raw)
    raw = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", raw)
    raw = re.sub(r"^\s{0,3}#{1,6}\s*", "", raw, flags=re.MULTILINE)
    raw = re.sub(r"^\s*[-*+]\s+", "", raw, flags=re.MULTILINE)
    return raw


def doc_stream_candidates(content: bytes) -> list[bytes]:
    """Read likely text-bearing streams from a legacy Word OLE document."""
    if olefile is None or not olefile.isOleFile(io.BytesIO(content)):
        return []
    streams: list[bytes] = []
    try:
        with olefile.OleFileIO(io.BytesIO(content)) as ole:
            preferred = [["WordDocument"], ["1Table"], ["0Table"]]
            for stream_name in preferred:
                if ole.exists(stream_name):
                    streams.append(ole.openstream(stream_name).read())
            for stream_name in ole.listdir(streams=True):
                lowered = "/".join(stream_name).lower()
                if any(part in lowered for part in ["worddocument", "table", "summaryinformation"]):
                    blob = ole.openstream(stream_name).read()
                    if blob not in streams:
                        streams.append(blob)
    except Exception:
        return streams
    return streams


def decode_text_best_effort(content: bytes) -> str:
    for encoding in ["utf-8-sig", "utf-8"]:
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue

    candidates: list[str] = []
    for encoding in ["gb18030", "utf-16le", "utf-16be", "big5"]:
        try:
            candidates.append(content.decode(encoding, errors="ignore"))
        except Exception:
            continue
    return max(candidates, key=text_quality_score, default="")


def extract_readable_text(content: bytes) -> str:
    decoded = decode_text_best_effort(content)
    decoded = decoded.replace("\x00", "\n")
    decoded = re.sub(r"[\x01-\x08\x0b\x0c\x0e-\x1f\x7f-\x9f]", " ", decoded)
    chunks = re.findall(
        r"[\u4e00-\u9fffA-Za-z0-9][\u4e00-\u9fffA-Za-z0-9，。！？；：、,.!?;:()\[\]《》“”\"'/@%+\-_\s]{1,}",
        decoded,
    )
    lines = []
    for chunk in chunks:
        line = re.sub(r"\s+", " ", chunk).strip(" \t\r\n-_")
        if len(line) >= 2 and text_quality_score(line) >= 1:
            lines.append(line)
    return "\n".join(dedupe_keep_order(lines))


def dedupe_keep_order(values: list[str]) -> list[str]:
    seen: set[str] = set()
    result: list[str] = []
    for value in values:
        key = value[:120]
        if key in seen:
            continue
        seen.add(key)
        result.append(value)
    return result


def text_quality_score(value: str) -> int:
    if not value:
        return 0
    chinese = len(re.findall(r"[\u4e00-\u9fff]", value))
    letters = len(re.findall(r"[A-Za-z]", value))
    digits = len(re.findall(r"\d", value))
    punctuation = len(re.findall(r"[，。！？；：、,.!?;:()\[\]《》“”\"']", value))
    controls = len(re.findall(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f-\x9f]", value))
    replacements = value.count("\ufffd")
    return chinese * 3 + letters + digits + punctuation - controls * 3 - replacements * 8


def is_probably_garbled(value: str) -> bool:
    clean = value.strip()
    if not clean:
        return True
    readable = len(re.findall(r"[\u4e00-\u9fffA-Za-z0-9]", clean))
    odd = len(re.findall(r"[^\u4e00-\u9fffA-Za-z0-9，。！？；：、,.!?;:()\[\]《》“”\"'/@%+\-_\s]", clean))
    if readable < 20:
        return True
    return odd / max(len(clean), 1) > 0.18


def normalize_text(value: str) -> str:
    text = unescape(value)
    text = text.replace("\u00a0", " ")
    text = text.replace("\x00", "")
    text = re.sub(r"[\x01-\x08\x0b\x0c\x0e-\x1f\x7f]", " ", text)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def summarize(value: str, limit: int = 180) -> str:
    clean = re.sub(r"\s+", " ", value).strip()
    if not clean:
        return "文件已解析，但没有提取到有效文本。"
    return clean[:limit] + ("..." if len(clean) > limit else "")


def plain_text_to_html(value: str) -> str:
    parts = []
    for line in str(value or "").splitlines():
        clean = line.strip()
        if not clean:
            parts.append("<p><br /></p>")
        elif re.match(r"^【.+】$", clean) or re.match(r"^(教育背景|实习经历|工作经历|项目经历|个人技能|竞赛经历|获奖经历|求职意向)$", clean):
            parts.append(f"<h3>{escape(clean)}</h3>")
        elif re.match(r"^[•·*\-]\s*", clean):
            bullet_text = re.sub(r"^[•·*\-]\s*", "", clean)
            parts.append(f"<p>• {escape(bullet_text)}</p>")
        else:
            parts.append(f"<p>{escape(clean)}</p>")
    return "\n".join(parts)


def file_title(filename: str, content: str) -> str:
    stem = Path(filename).stem.strip()
    first_line = next((line.strip() for line in content.splitlines() if line.strip()), "")
    if first_line and len(first_line) <= 40:
        return first_line
    return stem or "上传文件"


def title_from_html(content: bytes) -> str:
    raw = content.decode("utf-8", errors="ignore") or content.decode("gb18030", errors="ignore")
    soup = BeautifulSoup(raw, "html.parser")
    if soup.title and soup.title.string:
        return normalize_text(soup.title.string)[:120]
    heading = soup.find(["h1", "h2"])
    if heading:
        return normalize_text(heading.get_text(" "))[:120]
    return ""
